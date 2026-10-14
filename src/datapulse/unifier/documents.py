"""Read unstructured inputs - PDF, images, Word, Markdown/text, email, HTML, zip - into Documents."""
from __future__ import annotations

import email
import email.policy
import io as _io
import zipfile
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import List, Optional, Tuple

from datapulse.errors import DataPulseError
from datapulse.io import ARCHIVES, DOCUMENTS, INPUTS, SUPPORTED

MAX_DOCUMENTS = 50
MAX_TOTAL_BYTES = 25 * 1024 * 1024
MAX_TEXT_CHARS = 100_000
IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp",
               ".gif": "image/gif"}


@dataclass
class Document:
    name: str
    kind: str  # pdf | image | docx | text | email | html
    text: str = ""
    data: Optional[bytes] = None  # PDFs and images go to the AI as files, so it can read scans
    media_type: Optional[str] = None
    pages: Optional[int] = None
    size: int = 0
    truncated: bool = False


def _cap(doc: Document) -> Document:
    if len(doc.text) > MAX_TEXT_CHARS:
        doc.text = doc.text[:MAX_TEXT_CHARS]
        doc.truncated = True
    return doc


# ---------- per format ----------

def _pdf(name: str, data: bytes) -> Document:
    text, pages = "", None
    try:
        from pypdf import PdfReader

        reader = PdfReader(_io.BytesIO(data))
        pages = len(reader.pages)
        text = "\n\n".join(f"[page {n}]\n{(page.extract_text() or '').strip()}"
                           for n, page in enumerate(reader.pages, 1))
    except ImportError:
        pass  # the AI still reads the PDF itself; only the offline fallback loses its text
    except Exception as exc:
        raise DataPulseError(f"{name} is not a readable PDF: {exc}")
    return Document(name=name, kind="pdf", text=text, data=data, media_type="application/pdf", pages=pages,
                    size=len(data))


def _docx(name: str, data: bytes) -> Document:
    try:
        import docx
    except ImportError as exc:
        raise DataPulseError("Reading .docx files needs python-docx: pip install python-docx") from exc
    try:
        d = docx.Document(_io.BytesIO(data))
    except Exception as exc:
        raise DataPulseError(f"{name} is not a readable Word document: {exc}")
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        parts.append("\n".join("\t".join(cell.text.strip() for cell in row.cells) for row in table.rows))
    return Document(name=name, kind="docx", text="\n".join(parts), size=len(data))


class _TextFromHTML(HTMLParser):
    SKIP = {"script", "style", "noscript", "template", "svg", "head"}
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table",
             "ul", "ol", "header", "footer", "pre", "blockquote"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: List[str] = []
        self.title = ""
        self._skip = 0
        self._in_title = False

    def handle_starttag(self, tag, attrs):
        if tag == "title":
            self._in_title = True
        if tag in self.SKIP:
            self._skip += 1
        elif tag in ("td", "th"):
            self.parts.append("\t")
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag == "title":
            self._in_title = False
        if tag in self.SKIP and self._skip:
            self._skip -= 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        elif not self._skip:
            self.parts.append(data)

    def text(self) -> str:
        lines = []
        for line in "".join(self.parts).splitlines():
            cells = [" ".join(cell.split()) for cell in line.split("\t")]  # tabs separate table cells
            lines.append("\t".join(cell for cell in cells if cell))
        return "\n".join(line for line in lines if line)


def html_to_text(html: str) -> Tuple[str, str]:
    parser = _TextFromHTML()
    parser.feed(html)
    parser.close()
    return parser.title.strip(), parser.text()


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _html(name: str, data: bytes) -> Document:
    title, text = html_to_text(_decode(data))
    return Document(name=name, kind="html", text=(f"{title}\n\n{text}" if title else text), size=len(data))


def _email(name: str, data: bytes) -> List[Document]:
    msg = email.message_from_bytes(data, policy=email.policy.default)
    header = "\n".join(f"{h}: {msg[h]}" for h in ("Subject", "From", "To", "Cc", "Date") if msg[h])
    body = msg.get_body(preferencelist=("plain", "html"))
    text = ""
    if body is not None:
        content = body.get_content()
        text = html_to_text(content)[1] if body.get_content_subtype() == "html" else content
    docs = [Document(name=name, kind="email", text=f"{header}\n\n{text.strip()}", size=len(data))]
    for part in msg.iter_attachments():
        filename = part.get_filename()
        if not filename or Path(filename).suffix.lower() not in DOCUMENTS:
            continue
        payload = part.get_payload(decode=True) or b""
        docs.extend(_read_one(f"{name} / {filename}", payload, Path(filename).suffix))
    return docs


def _read_one(name: str, data: bytes, ext: str) -> List[Document]:
    ext = ext.lower()
    if ext == ".pdf":
        return [_cap(_pdf(name, data))]
    if ext in IMAGE_TYPES:
        return [Document(name=name, kind="image", data=data, media_type=IMAGE_TYPES[ext], size=len(data))]
    if ext == ".docx":
        return [_cap(_docx(name, data))]
    if ext in (".md", ".markdown", ".txt"):
        return [_cap(Document(name=name, kind="text", text=_decode(data), size=len(data)))]
    if ext in (".html", ".htm"):
        return [_cap(_html(name, data))]
    if ext == ".eml":
        return [_cap(d) for d in _email(name, data)]
    raise DataPulseError(f"Unsupported document type '{ext}' for {name}")


def full_text(path: Path, name: str) -> Optional[str]:
    """All the text of one PDF, Word, text or HTML file, without the AI prompt cap. None for other kinds."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in (".pdf", ".docx", ".md", ".markdown", ".txt", ".html", ".htm"):
        return None
    if path.stat().st_size > MAX_TOTAL_BYTES:
        raise DataPulseError(f"{name} is over the 25 MB limit for one upload or pull")
    data = path.read_bytes()
    if ext == ".pdf":
        return _pdf(name, data).text
    if ext == ".docx":
        return _docx(name, data).text
    if ext in (".html", ".htm"):
        return html_to_text(_decode(data))[1]
    return _decode(data)


def read_documents(files: List[Tuple[Path, str]]) -> Tuple[List[Document], List[dict]]:
    """files: (path on disk, name the user knows). Returns (documents, skipped [{name, reason}])."""
    docs: List[Document] = []
    skipped: List[dict] = []
    total = 0
    for path, name in files:
        data = Path(path).read_bytes()
        total += len(data)
        if total > MAX_TOTAL_BYTES:
            raise DataPulseError("Documents exceed the 25 MB limit for one upload or pull")
        try:
            docs.extend(_read_one(name, data, Path(path).suffix))
        except DataPulseError as exc:
            skipped.append({"name": name, "reason": str(exc)})
    if len(docs) > MAX_DOCUMENTS:
        skipped.extend({"name": d.name, "reason": f"over the {MAX_DOCUMENTS}-document limit"}
                       for d in docs[MAX_DOCUMENTS:])
        docs = docs[:MAX_DOCUMENTS]
    if not docs:
        reasons = "; ".join(f"{s['name']}: {s['reason']}" for s in skipped[:3])
        raise DataPulseError(f"No readable documents{f' ({reasons})' if reasons else ''}")
    return docs, skipped


def expand_zip(path: Path, dest: Path, label: str) -> Tuple[List[Tuple[Path, str]], List[dict]]:
    """Unpack a zip's supported files into dest. Guards against zip bombs and path tricks."""
    out: List[Tuple[Path, str]] = []
    skipped: List[dict] = []
    try:
        archive = zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        raise DataPulseError(f"{label} is not a valid zip file: {exc}")
    with archive:
        members = [m for m in archive.infolist() if not m.is_dir() and not Path(m.filename).name.startswith(".")
                   and "__MACOSX" not in m.filename]
        if sum(m.file_size for m in members) > MAX_TOTAL_BYTES:
            raise DataPulseError(f"{label} unpacks to more than 25 MB")
        for m in members:
            inner = Path(m.filename).name  # flatten folders; never write outside dest
            ext = Path(inner).suffix.lower()
            if ext not in INPUTS or ext in ARCHIVES:
                skipped.append({"name": f"{label} / {m.filename}", "reason": f"unsupported file type '{ext}'"})
                continue
            target = dest / f"{len(out):03d}-{inner}"
            target.write_bytes(archive.read(m))
            out.append((target, f"{label} / {m.filename}"))
    if not out:
        raise DataPulseError(f"{label} has no supported files")
    return out, skipped


def is_document(path_or_name) -> bool:
    return Path(str(path_or_name)).suffix.lower() in DOCUMENTS


def is_tabular(path_or_name) -> bool:
    return Path(str(path_or_name)).suffix.lower() in SUPPORTED
