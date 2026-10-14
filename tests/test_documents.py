"""The AI data unifier, extract path: document readers, limits, and the extraction agent."""
import io
import json
import zipfile
from email.message import EmailMessage
from pathlib import Path

import pytest

from datapulse.config import Config
from datapulse.errors import DataPulseError
from datapulse.unifier import documents as docs_mod
from datapulse.unifier import extract as extract_mod
from datapulse.unifier import unify_files
from datapulse.unifier.documents import Document, expand_zip, html_to_text, read_documents
from datapulse.unifier.extract import content_blocks, extract_all, propose_extraction_schema, records_schema
from datapulse.unifier.spec import UnifySpec
from fake_claude import FakeClaude, text


def make_pdf(line: str) -> bytes:
    """A minimal one-page PDF with real text, so pypdf can extract it."""
    stream = f"BT /F1 12 Tf 72 720 Td ({line}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % n + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def make_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph("Statement of work for Acme Corp")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Price"
    table.cell(1, 0).text, table.cell(1, 1).text = "Design", "1200"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_eml() -> bytes:
    msg = EmailMessage()
    msg["Subject"], msg["From"], msg["To"] = "Invoice INV-7", "billing@acme.com", "ap@us.com"
    msg.set_content("Please find the invoice attached. Total due: 1,250.00")
    msg.add_attachment(make_pdf("Invoice INV-7 total 1250"), maintype="application", subtype="pdf",
                       filename="inv-7.pdf")
    msg.add_attachment(b"MZ", maintype="application", subtype="octet-stream", filename="tool.exe")
    return msg.as_bytes()


HTML = b"""<html><head><title>Price list</title><style>.x{}</style></head>
<body><script>var secret = 1;</script><h1>Widgets</h1><table><tr><td>Bolt</td><td>0.25</td></tr></table></body></html>"""


def _write(tmp_path, name, data):
    p = tmp_path / name
    p.write_bytes(data if isinstance(data, bytes) else data.encode())
    return p, name


# ---------- readers ----------

def test_readers_cover_every_document_type(tmp_path):
    files = [_write(tmp_path, "a.pdf", make_pdf("Invoice 42 total 99.50")),
             _write(tmp_path, "b.docx", make_docx()),
             _write(tmp_path, "c.md", "# Notes\n\n- one"),
             _write(tmp_path, "d.txt", "plain text"),
             _write(tmp_path, "e.eml", make_eml()),
             _write(tmp_path, "f.html", HTML),
             _write(tmp_path, "g.png", b"\x89PNG fake")]
    docs, skipped = read_documents(files)
    by = {d.name: d for d in docs}
    assert "Invoice 42 total 99.50" in by["a.pdf"].text and by["a.pdf"].pages == 1 and by["a.pdf"].data
    assert "Acme Corp" in by["b.docx"].text and "Design\t1200" in by["b.docx"].text
    assert by["c.md"].kind == "text" and by["d.txt"].text == "plain text"
    assert by["e.eml"].text.startswith("Subject: Invoice INV-7") and "1,250.00" in by["e.eml"].text
    assert "INV-7 total 1250" in by["e.eml / inv-7.pdf"].text  # attachments become their own documents
    assert not any("tool.exe" in n for n in by)
    assert by["f.html"].text.startswith("Price list") and "Bolt\t0.25" in by["f.html"].text
    assert "secret" not in by["f.html"].text
    assert by["g.png"].kind == "image" and by["g.png"].media_type == "image/png" and skipped == []


def test_unreadable_documents_are_skipped_not_fatal(tmp_path):
    docs, skipped = read_documents([_write(tmp_path, "ok.txt", "hi"), _write(tmp_path, "bad.pdf", b"not a pdf")])
    assert [d.name for d in docs] == ["ok.txt"] and skipped[0]["name"] == "bad.pdf"
    with pytest.raises(DataPulseError, match="No readable documents"):
        read_documents([_write(tmp_path, "bad2.pdf", b"nope")])


def test_html_to_text_keeps_title_and_structure():
    title, body = html_to_text("<title>T</title><p>One</p><p>Two &amp; three</p>")
    assert title == "T" and body == "One\nTwo & three"


def test_document_limits(tmp_path, monkeypatch):
    files = [_write(tmp_path, f"{i}.txt", "x") for i in range(docs_mod.MAX_DOCUMENTS + 3)]
    docs, skipped = read_documents(files)
    assert len(docs) == docs_mod.MAX_DOCUMENTS and len(skipped) == 3
    monkeypatch.setattr(docs_mod, "MAX_TEXT_CHARS", 10)
    doc = read_documents([_write(tmp_path, "long.txt", "y" * 50)])[0][0]
    assert doc.truncated and len(doc.text) == 10
    monkeypatch.setattr(docs_mod, "MAX_TOTAL_BYTES", 20)
    with pytest.raises(DataPulseError, match="25 MB"):
        read_documents([_write(tmp_path, "big.txt", "z" * 30)])


def test_zip_is_flattened_filtered_and_bounded(tmp_path, monkeypatch):
    archive = tmp_path / "bundle.zip"
    with zipfile.ZipFile(archive, "w") as z:
        z.writestr("docs/a.txt", "alpha")
        z.writestr("../../escape.md", "beta")
        z.writestr("__MACOSX/._a.txt", "junk")
        z.writestr("tool.exe", "MZ")
    dest = tmp_path / "out"
    dest.mkdir()
    out, skipped = expand_zip(archive, dest, "bundle.zip")
    assert [label for _, label in out] == ["bundle.zip / docs/a.txt", "bundle.zip / ../../escape.md"]
    assert all(p.parent == dest for p, _ in out)  # nothing written outside dest
    assert [s["name"] for s in skipped] == ["bundle.zip / tool.exe"]
    monkeypatch.setattr(docs_mod, "MAX_TOTAL_BYTES", 3)
    with pytest.raises(DataPulseError, match="unpacks to more than"):
        expand_zip(archive, dest, "bundle.zip")


def test_pdfs_and_images_go_to_the_ai_as_files():
    pdf = content_blocks(Document(name="a.pdf", kind="pdf", data=b"%PDF", media_type="application/pdf", pages=2))
    assert pdf[1]["type"] == "document" and pdf[1]["source"]["media_type"] == "application/pdf"
    img = content_blocks(Document(name="scan.jpg", kind="image", data=b"jpg", media_type="image/jpeg"))
    assert img[1]["type"] == "image" and img[1]["source"]["media_type"] == "image/jpeg"
    txt = content_blocks(Document(name="n.txt", kind="text", text="hello"))
    assert len(txt) == 1 and txt[0]["text"].endswith("hello")


# ---------- extraction agent ----------

INVOICE_SCHEMA = {"description": "Invoice line items", "record_mode": "many_per_document",
                  "record_description": "One invoice line", "fields": [
                      {"name": "invoice_number", "type": "string", "description": "Invoice id"},
                      {"name": "invoice_date", "type": "date", "description": "Date issued"},
                      {"name": "item", "type": "string", "description": "Line item"},
                      {"name": "amount", "type": "number", "description": "Line amount"}]}


def _lines(*rows):
    return text(json.dumps({"records": [dict(zip(["invoice_number", "invoice_date", "item", "amount", "_page",
                                                  "_confidence"], r)) for r in rows]}))


@pytest.fixture
def serial(monkeypatch):
    monkeypatch.setattr(extract_mod, "CONCURRENCY", 1)


def test_schema_proposal_and_many_records_per_document(serial):
    docs = [Document(name="a.txt", kind="text", text="INV-1 ..."), Document(name="b.txt", kind="text", text="INV-2")]
    fake = FakeClaude([text(json.dumps(INVOICE_SCHEMA)),
                       _lines(["INV-1", "2026-03-07", "Bolts", 12.5, 1, "high"],
                              ["INV-1", "2026-03-07", "Nuts", "1,000", 1, "medium"]),
                       _lines(["INV-2", "early March", "Washers", 3, None, "low"])])
    spec = propose_extraction_schema(docs, None, Config(), client=fake)
    assert spec.mode == "extract" and spec.record_mode == "many_per_document"
    rows, stats = extract_all(docs, spec, Config(), client=fake)
    assert rows[0] == {"invoice_number": "INV-1", "invoice_date": "2026-03-07", "item": "Bolts", "amount": 12.5,
                       "_document": "a.txt", "_page": 1, "_confidence": "high"}
    assert rows[1]["amount"] is None  # "1,000" is not a plain number; never silently fixed
    assert rows[2]["_document"] == "b.txt" and rows[2]["_page"] is None and rows[2]["invoice_date"] is None
    assert stats["parse_errors"] == {"amount": 1, "invoice_date": 1} and stats["low_confidence"] == 1
    schema = fake.calls[1]["output_config"]["format"]["schema"]["properties"]["records"]["items"]
    assert set(schema["required"]) == {"invoice_number", "invoice_date", "item", "amount", "_page", "_confidence"}


def test_one_per_document_keeps_one_record_and_reports_failures(serial):
    spec = UnifySpec.from_dict({"mode": "extract", "record_mode": "one_per_document", "fields": [
        {"name": "subject", "type": "string"}, {"name": "total", "type": "number"}]})
    docs = [Document(name=n, kind="text", text="x") for n in ("a.txt", "b.txt", "c.txt")]
    two = text(json.dumps({"records": [{"subject": "A", "total": 1, "_page": None, "_confidence": "high"},
                                       {"subject": "extra", "total": 2, "_page": None, "_confidence": "high"}]}))
    fake = FakeClaude([two, text("not json"), text(json.dumps({"records": []}))])
    rows, stats = extract_all(docs, spec, Config(), client=fake)
    assert [r["subject"] for r in rows] == ["A"]
    assert stats["failed_documents"][0]["name"] == "b.txt" and stats["documents_without_records"] == ["c.txt"]
    assert "exactly one record" in fake.calls[0]["messages"][0]["content"][0]["text"]


def test_records_schema_rejects_extra_keys():
    spec = UnifySpec.from_dict({**INVOICE_SCHEMA, "mode": "extract"})
    item = records_schema(spec)["properties"]["records"]["items"]
    assert item["additionalProperties"] is False and item["properties"]["amount"]["type"] == ["number", "null"]


def test_without_a_key_each_document_becomes_one_record(tmp_path):
    files = [_write(tmp_path, "notes.md", "# Kickoff\n\nWe agreed on scope."), _write(tmp_path, "m.eml", make_eml())]
    out = unify_files(files, None, Config(), use_ai=False)
    assert out.mode == "extract" and out.origin == "heuristic"
    titles = {r["_document"]: r["title"] for r in out.rows}
    assert titles["notes.md"] == "# Kickoff" and titles["m.eml"] == "Invoice INV-7"
    assert out.rows[0]["word_count"] == 6 and "No AI key" in out.stats["note"]


def test_a_pdf_of_the_orders_table_is_the_orders_table(tmp_path):
    from datapulse.baseline import build_baseline
    from datapulse.checks import run_checks
    from datapulse.config import load_config
    from datapulse.io import load_dataframe
    from datapulse.sources.anyfile import frame_from_path
    from datapulse.sources.pdfdoc import table_pdf

    orders_csv = Path(__file__).resolve().parents[1] / "examples" / "data" / "orders.csv"
    orders = load_dataframe(orders_csv)
    pdf = tmp_path / "orders.pdf"
    pdf.write_bytes(table_pdf(orders))
    frame = frame_from_path(pdf, "orders.pdf")
    assert list(frame.columns) == list(orders.columns)
    assert len(frame) == len(orders)
    assert frame["order_id"].tolist() == orders["order_id"].tolist()
    assert frame["customer_id"].tolist() == orders["customer_id"].tolist()
    assert frame["status"].tolist() == orders["status"].tolist()
    assert frame["country"].tolist() == orders["country"].tolist()
    assert frame["amount"].tolist() == orders["amount"].tolist()

    cfg = load_config(orders_csv.parents[1] / "datapulse.yml")
    report = run_checks(build_baseline(orders, orders_csv, 1, cfg.max_categories), frame, cfg,
                        data_file="orders.pdf", reference=orders)
    assert not report.failed, [(i.severity, i.title) for i in report.issues]


@pytest.mark.parametrize("ext", [".pdf", ".docx", ".txt", ".html"])
def test_a_long_table_document_keeps_every_row(tmp_path, ext):
    import pandas as pd

    from datapulse.io import load_dataframe
    from datapulse.sources.anyfile import frame_from_path
    from datapulse.sources.pdfdoc import table_pdf
    from datapulse.unifier.documents import MAX_TEXT_CHARS

    orders = load_dataframe(Path(__file__).resolve().parents[1] / "examples" / "data" / "orders.csv")
    long = pd.concat([orders] * 6, ignore_index=True)
    assert len(long.to_csv(index=False)) > MAX_TEXT_CHARS
    path = tmp_path / f"long{ext}"
    cells = [list(long.columns)] + long.astype(str).values.tolist()
    if ext == ".pdf":
        path.write_bytes(table_pdf(long))
    elif ext == ".docx":
        docx = pytest.importorskip("docx")
        doc = docx.Document()
        table = doc.add_table(rows=len(cells), cols=len(long.columns))
        for row, values in zip(table.rows, cells):
            for cell, value in zip(row.cells, values):
                cell.text = value
        doc.save(path)
    elif ext == ".txt":
        path.write_text(long.to_csv(index=False))
    else:
        rows = "".join("<tr>" + "".join(f"<td>{v}</td>" for v in values) + "</tr>" for values in cells)
        path.write_text(f"<html><head><title>Orders</title></head><body><table>{rows}</table></body></html>")
    frame = frame_from_path(path, path.name)
    assert len(frame) == len(long)
    assert frame.astype(str).equals(long.astype(str))


def test_documents_can_be_added_as_a_source(tmp_path):
    from datapulse.sources.anyfile import frame_from_bytes, frame_from_path
    from datapulse.sources.link import read_bytes

    pdf = tmp_path / "invoice.pdf"
    pdf.write_bytes(make_pdf("Invoice 42 total 99.50"))
    frame = frame_from_path(pdf)
    assert len(frame) == 1 and "Invoice 42 total 99.50" in frame.loc[0, "text"]

    docx = tmp_path / "sow.docx"
    docx.write_bytes(make_docx())
    frame = frame_from_path(docx)
    assert "Acme Corp" in frame.loc[0, "text"] and frame.loc[0, "_document"] == "sow.docx"

    eml = tmp_path / "mail.eml"
    eml.write_bytes(make_eml())
    frame = frame_from_path(eml)
    assert "Invoice INV-7" in set(frame["title"]) or "Invoice INV-7" in " ".join(frame["text"])
    assert any("inv-7.pdf" in name for name in frame["_document"])

    linked = read_bytes(make_pdf("Hello from a link"), "https://cdn.example.com/note.pdf", "application/pdf")
    assert "Hello from a link" in linked.loc[0, "text"]
    linked = frame_from_bytes(make_docx(), "brief.docx")
    assert "Acme Corp" in linked.loc[0, "text"]


def test_web_page_source_saves_html(tmp_path, monkeypatch):
    httpx = pytest.importorskip("httpx")
    from datapulse.sources import http as http_mod
    from datapulse.sources.web_page import WebPageSource

    monkeypatch.setattr(http_mod, "_TRANSPORT", httpx.MockTransport(
        lambda request: httpx.Response(200, content=HTML, headers={"content-type": "text/html"})))
    monkeypatch.setattr(http_mod, "_check_host", lambda host: None)
    path = WebPageSource().fetch({"url": "https://shop.example/prices"}, {}, tmp_path)
    assert path.suffix == ".html" and b"Price list" in path.read_bytes() and b"Bolt" in path.read_bytes()
